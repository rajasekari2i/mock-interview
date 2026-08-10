"""Bounded, signature-aware JD document extraction."""

from __future__ import annotations

from dataclasses import dataclass
from io import BytesIO
from pathlib import PurePosixPath
from zipfile import BadZipFile, ZipFile

from docx import Document
from pypdf import PdfReader

from app.auth.errors import AuthError, ErrorCode

MAX_UPLOAD_BYTES = 5 * 1024 * 1024
MAX_ARCHIVE_ENTRIES = 500
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 20 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100


@dataclass(frozen=True)
class ParsedDocument:
    source_format: str
    text: str
    original_filename: str


def _reject() -> AuthError:
    return AuthError(ErrorCode.UNSUPPORTED_DOCUMENT)


def _normalized_text(value: str) -> str:
    normalized = " ".join(value.split())
    if not normalized:
        raise _reject()
    return normalized


def _safe_basename(filename: str) -> str:
    return PurePosixPath(filename.replace("\\", "/")).name or "document"


def _extract_pdf(payload: bytes) -> str:
    if not payload.startswith(b"%PDF-"):
        raise _reject()
    try:
        reader = PdfReader(BytesIO(payload))
        if reader.is_encrypted:
            raise _reject()
        return _normalized_text("\n".join(page.extract_text() or "" for page in reader.pages))
    except Exception as error:
        raise _reject() from error


def _validate_docx_archive(payload: bytes) -> None:
    try:
        with ZipFile(BytesIO(payload)) as archive:
            infos = archive.infolist()
            if not infos or len(infos) > MAX_ARCHIVE_ENTRIES:
                raise _reject()
            total = 0
            names = {info.filename for info in infos}
            for info in infos:
                path = PurePosixPath(info.filename)
                if path.is_absolute() or ".." in path.parts or info.flag_bits & 0x1:
                    raise _reject()
                total += info.file_size
                if total > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
                    raise _reject()
                if (
                    info.file_size > 0
                    and info.file_size > max(1, info.compress_size) * MAX_COMPRESSION_RATIO
                ):
                    raise _reject()
            if "[Content_Types].xml" not in names or "word/document.xml" not in names:
                raise _reject()
            content_types = archive.read("[Content_Types].xml")
            if b"macroEnabled" in content_types or any(
                name.casefold().endswith("vbaproject.bin") for name in names
            ):
                raise _reject()
    except (BadZipFile, KeyError, OSError, ValueError) as error:
        raise _reject() from error


def _extract_docx(payload: bytes) -> str:
    _validate_docx_archive(payload)
    try:
        document = Document(BytesIO(payload))
        return _normalized_text("\n".join(paragraph.text for paragraph in document.paragraphs))
    except Exception as error:
        raise _reject() from error


def _extract_txt(payload: bytes) -> str:
    if b"\x00" in payload:
        raise _reject()
    try:
        return _normalized_text(payload.decode("utf-8", errors="strict"))
    except UnicodeDecodeError as error:
        raise _reject() from error


def parse_document(*, filename: str, content_type: str, payload: bytes) -> ParsedDocument:
    if len(payload) > MAX_UPLOAD_BYTES:
        raise AuthError(ErrorCode.UPLOAD_TOO_LARGE)
    suffix = PurePosixPath(filename.casefold()).suffix
    handlers = {
        (".pdf", "application/pdf"): ("PDF", _extract_pdf),
        (
            ".docx",
            "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        ): ("DOCX", _extract_docx),
        (".txt", "text/plain"): ("TXT", _extract_txt),
    }
    selected = handlers.get((suffix, content_type.casefold()))
    if selected is None:
        raise _reject()
    source_format, extractor = selected
    return ParsedDocument(
        source_format=source_format,
        text=extractor(payload),
        original_filename=_safe_basename(filename),
    )
